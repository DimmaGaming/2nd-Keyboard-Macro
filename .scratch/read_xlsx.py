import openpyxl

wb = openpyxl.load_workbook('d:/Dimma/Coding Workflow/AutoHotKey Scripts/docs/Macro Keyboard Registry NEW.xlsx')
with open('d:/Dimma/Coding Workflow/AutoHotKey Scripts/.scratch/spreadsheet_dump.txt', 'w', encoding='utf-8') as f:
    for sheetname in wb.sheetnames:
        f.write(f"=== Sheet: {sheetname} ===\n")
        sheet = wb[sheetname]
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None and str(cell).strip() != "" for cell in row):
                f.write(" | ".join(str(c) if c is not None else "" for c in row) + "\n")
