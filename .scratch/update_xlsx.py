import openpyxl
import sys

def update_spreadsheet():
    filepath = 'd:/Dimma/Coding Workflow/AutoHotKey Scripts/docs/Macro Keyboard Registry NEW.xlsx'
    wb = openpyxl.load_workbook(filepath)
    sheet = wb['Macro Keyboard Registry']
    
    # Define mapping from F-Key + Modifiers -> New Action
    updates = {
        ('F13', 'Ctrl+Shift+Alt'): {
            'Status': 'USED',
            'Action': 'Normal: AI Apps (ChatGPT/Gemini)\n+CapsLock: Slide Left In\n+Caps+Shift: Slide Left Out'
        },
        ('F14', 'Ctrl+Shift+Alt'): {
            'Status': 'USED',
            'Action': 'Normal: Premiere Pro\n+CapsLock: Slide Up In\n+Caps+Shift: Slide Up Out'
        },
        ('F15', 'Ctrl+Shift+Alt'): {
            'Status': 'USED',
            'Action': 'Normal: Brave/Docs Cycler\n+CapsLock: Slide Right In\n+Caps+Shift: Slide Right Out'
        },
        ('F16', 'Ctrl+Shift+Alt'): {
            'Status': 'USED',
            'Action': 'Normal: OneCommander\n+CapsLock: Slide Down In\n+Caps+Shift: Slide Down Out'
        },
        ('F19', 'Shift'): {
            'Status': 'USED',
            'Action': 'Enter Editing Mode'
        },
        ('F19', 'Ctrl+Shift'): {
            'Status': 'USED',
            'Action': 'Exit Editing Mode'
        }
    }
    
    # Headers should be at row 1
    # Col 1: F-Key (A)
    # Col 2: Modifiers (B)
    # Col 5: Status (E)
    # Col 6: Assigned Action (F)
    
    changes_made = 0
    for row in range(2, sheet.max_row + 1):
        fkey = str(sheet.cell(row=row, column=1).value).strip()
        mod = str(sheet.cell(row=row, column=2).value).strip()
        
        if (fkey, mod) in updates:
            upd = updates[(fkey, mod)]
            sheet.cell(row=row, column=5).value = upd['Status']
            sheet.cell(row=row, column=6).value = upd['Action']
            changes_made += 1
            print(f"Updated {fkey} ({mod}) -> {upd['Action'].replace(chr(10), ' | ')}")
            
    wb.save(filepath)
    print(f"Successfully updated {changes_made} rows in {filepath}")

if __name__ == '__main__':
    update_spreadsheet()
